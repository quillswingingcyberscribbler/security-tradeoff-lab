#!/usr/bin/env python3
"""Read-only regression and boundary checks for Opportunity-Cost Security v6.2.

Reads source formulas and evaluates in-memory changes with a purpose-limited
interpreter. Does not save workbooks or automate Microsoft Excel. Independent
arithmetic checks use input values rather than translating workbook formulas.
"""
import argparse
import hashlib
import json
import math
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
from excel_eval import Evaluator, num

AS_OF = 46318  # 2026-10-23, fixed for reproducibility.
CASE_TABLES = ['Cases', 'Current', 'Proposed', 'Review & Decision',
               'Challenge Review', 'Actuals', 'Classification']


class OpportunityEvaluator(Evaluator):
    def eval(self, node, sheet):
        if node[0] == 'fn' and node[1] == 'ROUND':
            x = num(self.eval(node[2][0], sheet))
            n = int(num(self.eval(node[2][1], sheet)))
            return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))
        if node[0] == 'fn' and node[1] == 'HYPERLINK':
            return self.eval(node[2][-1], sheet)
        return super().eval(node, sheet)


def same(a, b):
    if a in (None, '') and b in (None, ''):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-8)
    return a == b


def approved_fixture(row=18):
    """Synthetic Live forecast with a complete, positive approval record."""
    records = {
        'Cases': {'B': 'Live', 'C': 'Synthetic process comparison', 'D': 'Synthetic scope',
                  'E': 'Test owner', 'F': 'Specified protection outcome', 'G': 'Replace',
                  'H': 'Full post-change profile', 'I': 'Synthetic assumptions'},
        'Current': {'B': 'Staff', 'C': 1000, 'D': 100, 'E': 10000, 'F': 'Synthetic baseline'},
        'Proposed': {'B': 'Staff', 'C': 200, 'D': 40, 'E': 5000, 'F': 3000, 'G': 'Synthetic forecast'},
        'Review & Decision': {'B': 'Evidence-backed', 'C': 'Equivalent or better',
                              'D': 'Synthetic comparative evidence', 'E': 'Approved',
                              'F': 'Test approver', 'G': AS_OF - 40, 'H': 'Validated',
                              'I': AS_OF - 30, 'J': 0.5, 'K': 0.75},
        'Challenge Review': {'B': 'Test challenger', 'C': AS_OF - 41,
                             **{c: 'Yes' for c in 'DEFGHI'}},
    }
    result = {(s, c + str(row)): v for s, cols in records.items() for c, v in cols.items()}
    result[('Settings', 'B7')] = AS_OF
    return result


def measured_fixture(row=18):
    result = approved_fixture(row)
    result[('Review & Decision', 'E' + str(row))] = 'Implemented'
    result[('Review & Decision', 'H' + str(row))] = 'Realized'
    actuals = {'B': 'Annualized measurement', 'C': AS_OF - 29, 'D': AS_OF - 1,
               'E': 600, 'F': 40, 'G': 2000, 'H': 35000, 'I': 'Synthetic measured evidence',
               'J': 'Test verifier', 'K': AS_OF, 'L': 'Annual equivalent; matched scope and fixed rates'}
    result.update({('Actuals', c + str(row)): v for c, v in actuals.items()})
    return result


def demand_fixture(row=36, linked=False):
    data = {'B': 'Live', 'C': 'Synthetic additional activity', 'D': 'Approved',
            'E': 100, 'F': 20, 'G': 1000, 'H': 2000,
            'I': 'C012' if linked else None, 'J': 'Yes' if linked else 'No',
            'K': 'Synthetic cost, overlap and resource-fit record'}
    return {('Reinvestment', c + str(row)): v for c, v in data.items()}


def run(path, recalculated=None):
    ev = OpportunityEvaluator(path)
    wb = load_workbook(recalculated or path, read_only=True, data_only=True)
    cache = {s.title: {c.coordinate: c.value for row in s for c in row if c.value is not None} for s in wb}
    wb.close()
    matched, mismatches, errors = 0, [], []
    for sheet, cells in ev.cells.items():
        for ref, (value, _) in cells.items():
            if not (isinstance(value, str) and value.startswith('=')):
                continue
            try:
                got, want = ev.cell(sheet, ref), cache[sheet].get(ref)
                if same(got, want):
                    matched += 1
                else:
                    mismatches.append({'sheet': sheet, 'cell': ref, 'expected': want, 'actual': got})
            except Exception as exc:
                errors.append({'sheet': sheet, 'cell': ref, 'error': str(exc)})
    structural = []
    expected_ids = [f'C{i:03d}' for i in range(1, 51)]
    for sheet in CASE_TABLES + ['Checks']:
        ids = [ev.raw(sheet, 'A' + str(r)) for r in range(7, 57)]
        structural.append({'check': sheet + ' has each prepared ID exactly once', 'pass': sorted(ids) == expected_ids})
    for sheet, col in [('Review & Decision', 'M'), ('Checks', 'J')]:
        structural.append({'check': sheet + ' has 50 readiness formulas',
                           'pass': all(isinstance(ev.raw(sheet, col + str(r)), str)
                                       and ev.raw(sheet, col + str(r)).startswith('=') for r in range(7, 57))})
    references = []
    rates = {g: ev.raw('Settings', c) for g, c in [('Staff', 'B9'), ('Management', 'B10'), ('Specialist', 'B11')]}
    sums = [0.0] * 5
    for row in range(7, 18):
        cur = {c: ev.raw('Current', c + str(row)) for c in 'BCDE'}
        pro = {c: ev.raw('Proposed', c + str(row)) for c in 'BCDEF'}
        ch = 0 if cur['B'] == 'Specialist' else cur['C']
        ph = 0 if pro['B'] == 'Specialist' else pro['C']
        ct = cur['D'] + (cur['C'] if cur['B'] == 'Specialist' else 0)
        pt = pro['D'] + (pro['C'] if pro['B'] == 'Specialist' else 0)
        cc = cur['C'] * rates[cur['B']] + cur['D'] * rates['Specialist'] + cur['E']
        pc = pro['C'] * rates[pro['B']] + pro['D'] * rates['Specialist'] + pro['E']
        expected = {'N': ch, 'O': ph, 'Q': ct, 'R': pt, 'W': cc, 'X': pc,
                    'P': ch - ph, 'S': ct - pt, 'V': cur['E'] - pro['E'], 'Y': cc - pc}
        for col, want in expected.items():
            got = ev.cell('Checks', col + str(row))
            references.append({'cell': 'Checks!' + col + str(row), 'expected': want, 'actual': got, 'pass': same(got, want)})
        for i, value in enumerate([cc, ch - ph, ct - pt, cur['E'] - pro['E'], cc - pc]):
            sums[i] += value
    for row, want in zip(range(39, 44), sums):
        got = ev.cell('Start', 'H' + str(row))
        references.append({'cell': 'Start!H' + str(row), 'expected': want, 'actual': got, 'pass': same(got, want)})
    scenarios = []

    def case(name, overrides, checks, topic, finding=None):
        ev.reset(overrides)
        results = []
        for sheet, ref, expected in checks:
            try:
                actual = ev.cell(sheet, ref)
            except Exception as exc:
                actual = {'evaluation_error': str(exc)}
            results.append({'sheet': sheet, 'cell': ref, 'expected': expected,
                            'actual': actual, 'matches': same(actual, expected)})
        scenarios.append({'name': name, 'topic': topic,
                          'kind': 'known-limitation reproduction' if finding else 'regression / boundary check',
                          'finding': finding, 'results': results})

    base, actual = approved_fixture(), measured_fixture()
    rd, ck, re = 'Review & Decision', 'Checks', 'Reinvestment'
    no_plan = [(ck, 'M18', 0), (ck, 'AA18', 0), (ck, 'AB18', 0), (ck, 'AC18', 0)]
    case('Baseline examples never supply the Live plan', {}, [('Start', 'A9', 0), ('Start', 'E9', 0),
         ('Start', 'I9', 0), ('Start', 'H39', 1423050), ('Start', 'H43', 899050), (re, 'E11', 0)], 'baseline')
    case('Complete approved forecast supplies usable hours and cash', base,
         [(ck, 'K18', 1), (ck, 'L18', 0), (ck, 'M18', 1), (ck, 'AA18', 400), (ck, 'AB18', 45),
          (ck, 'AC18', 5000), (ck, 'AD18', 3000), (ck, 'Y18', 57500), (rd, 'M18', 'Approved forecast eligible')], 'positive approval')
    case('Modeled Live potential is visible but not allocatable', {**base, (rd, 'H18'): 'Modeled'},
         no_plan + [('Start', 'E31', 800), (rd, 'M18', 'Modeled only — not allocatable')], 'stage gate')
    case('Approved Example forecast cannot enter planning', {**base, ('Cases', 'B18'): 'Example'},
         no_plan + [(ck, 'K18', 0), (rd, 'M18', 'Example only — never in live planning')], 'OC-04 example isolation')
    case('Implemented Example actuals cannot enter planning', {**actual, ('Cases', 'B18'): 'Example'},
         no_plan + [(ck, 'L18', 0)], 'OC-04 example isolation')
    for sheet, cells in [('Current', 'BCDEF'), ('Proposed', 'BCDEFG')]:
        for col in cells:
            case(f'{sheet} missing {col} blocks planning', {**base, (sheet, col + '18'): None},
                 no_plan + [(ck, 'AH18', 0)], 'OC-03 missing inputs')
    for sheet, cells in [('Current', 'CDE'), ('Proposed', 'CDEF')]:
        for col in cells:
            case(f'{sheet} confirmed zero {col} is accepted', {**base, (sheet, col + '18'): 0},
                 [(ck, 'AH18', 1), (ck, 'K18', 1)], 'OC-03 zero inputs')
    for sheet, cell, value in [('Current', 'C18', -1), ('Current', 'D18', '100'), ('Current', 'B18', 'Unknown'),
                               ('Proposed', 'F18', -1), ('Settings', 'B9', -1), ('Settings', 'B10', None), ('Settings', 'B11', '75')]:
        case(f'Invalid {sheet}!{cell} is excluded', {**base, (sheet, cell): value}, no_plan, 'numeric / label validation')
    for decision in ['Open', 'Pilot approved', 'Rejected', 'Unknown']:
        case(f'Human decision {decision} cannot supply planning', {**base, (rd, 'E18'): decision}, no_plan, 'positive approval')
    for confidence in ['Hypothesis', 'Plausible', None]:
        case(f'Local confidence {confidence!r} cannot supply planning', {**base, (rd, 'B18'): confidence}, no_plan, 'positive approval')
    for cell in ['D18', 'F18', 'G18', 'I18', 'J18', 'K18']:
        case(f'Missing approval field {cell} holds planning', {**base, (rd, cell): None}, no_plan, 'positive approval')
    for cell in ['D18', 'F18']:
        case(f'Whitespace approval field {cell} is missing', {**base, (rd, cell): '   '}, no_plan, 'input completeness')
    case('Accepted protection reduction needs a reference', {**base, (rd, 'C18'): 'Accepted reduction'}, no_plan, 'protection gate')
    case('Documented accepted protection reduction can be approved',
         {**base, (rd, 'C18'): 'Accepted reduction', (rd, 'L18'): 'Synthetic risk acceptance'}, [(ck, 'K18', 1)], 'protection gate')
    case('Not-assessed protection remains excluded', {**base, (rd, 'C18'): 'Not assessed'}, no_plan, 'protection gate')
    for cell, value in [('D18', 'No'), ('E18', 'Unchecked'), ('B18', None), ('C18', None), ('J18', 'Unresolved concern')]:
        case(f'Challenge field {cell} prevents clearance', {**base, ('Challenge Review', cell): value}, no_plan, 'challenge gate')
    for case_id in [12, 13, 14, 15, 30, 50]:
        row = case_id + 6
        case(f'Prepared C{case_id:03d} blocker follows its key', {**approved_fixture(row), ('Challenge Review', 'D' + str(row)): 'No'},
             [(ck, 'G' + str(row), 'Blocked'), (ck, 'M' + str(row), 0)], 'OC-02 spare-row key')
    for sheet in CASE_TABLES:
        case(f'Duplicate ID in {sheet} invalidates case', {**base, (sheet, 'A19'): 'C012'},
             no_plan + [(ck, 'AG18', 0), (rd, 'M18', 'Fix duplicate / missing case ID')], 'key integrity')
    case('Missing review ID invalidates case', {**base, (rd, 'A18'): None}, no_plan + [(ck, 'AG18', 0)], 'key integrity')
    moved = dict(base)
    for col in 'ABCDEF':
        moved[('Current', col + '19')] = base.get(('Current', col + '18'), ev.raw('Current', col + '18'))
        moved[('Current', col + '18')] = ev.raw('Current', col + '19')
    case('Current record moved with its ID still joins correctly', moved,
         [(ck, 'AG18', 1), (ck, 'K18', 1), (ck, 'AA18', 400), (ck, 'AB18', 45)], 'key-based joins')
    reviews_moved = dict(moved)
    for sheet, destination, columns in [(rd, 20, 'ABCDEFGHIJKL'), ('Challenge Review', 21, 'ABCDEFGHIJ')]:
        for col in columns:
            reviews_moved[(sheet, col + str(destination))] = base.get((sheet, col + '18'), ev.cells[sheet].get(col + '18', (None, 'n'))[0])
            reviews_moved[(sheet, col + '18')] = ev.cells[sheet].get(col + str(destination), (None, 'n'))[0]
    case('Independent input-table reordering retains approval by key', reviews_moved,
         [(ck, 'AG18', 1), (ck, 'G18', 'Clear'), (ck, 'K18', 1), (ck, 'AA18', 400)], 'key-based joins')
    case('Blocker in reordered challenge record still blocks its case',
         {**reviews_moved, ('Challenge Review', 'D21'): 'No'},
         [(ck, 'G18', 'Blocked'), (ck, 'M18', 0)], 'key-based joins')
    for sheet, cell, value in [(rd, 'G18', AS_OF + 1), (rd, 'I18', AS_OF + 1), (rd, 'I18', AS_OF - 41),
                               ('Challenge Review', 'C18', AS_OF + 1), ('Settings', 'B7', None),
                               ('Settings', 'B7', 'not a date'), ('Settings', 'B8', 'Unknown')]:
        case(f'Invalid or unavailable {sheet}!{cell} holds planning', {**base, (sheet, cell): value}, no_plan, 'dates / settings')
    for cell, value in [('J18', -0.1), ('K18', 1.1), ('J18', '50%')]:
        case(f'Invalid usable share {cell} holds planning', {**base, (rd, cell): value}, no_plan, 'usable capacity')
    case('Zero usable shares do not turn time into cash', {**base, (rd, 'J18'): 0, (rd, 'K18'): 0},
         [(ck, 'M18', 1), (ck, 'AA18', 0), (ck, 'AB18', 0), (ck, 'AC18', 5000)], 'usable capacity')
    case('Extra forecast work counts fully despite zero usable shares',
         {**base, ('Proposed', 'C18'): 1200, ('Proposed', 'D18'): 150, (rd, 'J18'): 0, (rd, 'K18'): 0},
         [(ck, 'AA18', -200), (ck, 'AB18', -50)], 'usable capacity')
    case('Retain with changed annual profile is incomplete', {**base, ('Cases', 'G18'): 'Retain'},
         no_plan + [(ck, 'F18', 0)], 'proposal semantics')
    retain = {**base, ('Cases', 'G18'): 'Retain', ('Proposed', 'C18'): 1000,
              ('Proposed', 'D18'): 100, ('Proposed', 'E18'): 10000, ('Proposed', 'F18'): 0}
    case('Unchanged Retain releases no resource', retain,
         no_plan + [(ck, 'AH18', 1), (ck, 'P18', 0), (ck, 'Y18', 0)], 'proposal semantics')
    case('Specialist execution counts as specialist capacity', {**base, ('Current', 'B18'): 'Specialist'},
         [(ck, 'N18', 0), (ck, 'Q18', 1100), (ck, 'P18', -200), (ck, 'S18', 1060)], 'resource classification')
    case('Realized without actuals is excluded and visibly flagged', {**base, (rd, 'H18'): 'Realized', (rd, 'E18'): 'Implemented'},
         no_plan + [(ck, 'AF18', 'Missing measured evidence'), (rd, 'M18', 'Add measurements, period and evidence'),
                    ('Start', 'A12', 'Action required: one or more live Realized cases lack complete measurement evidence.')], 'OC-01 readiness')
    case('Annualized actuals supply only entered annual equivalents', actual,
         [(ck, 'K18', 0), (ck, 'L18', 1), (ck, 'M18', 1), (ck, 'AA18', 300), (ck, 'AB18', 30),
          (ck, 'AC18', 2000), (ck, 'AD18', 0), (ck, 'AO18', 37000), ('Actuals', 'M18', 'Ready'),
          (rd, 'M18', 'Measured result eligible')], 'OC-04 actuals')
    for col in 'BCDEFGHIJKL':
        case(f'Missing Actuals {col} blocks realized planning', {**actual, ('Actuals', col + '18'): None},
             no_plan + [(ck, 'I18', 0), (ck, 'L18', 0)], 'OC-04 actuals completeness')
    for cell in ['I18', 'J18', 'L18']:
        case(f'Whitespace Actuals {cell} is incomplete', {**actual, ('Actuals', cell): '   '}, no_plan, 'OC-04 actuals provenance')
    case('Short period cannot be labelled full-year actual', {**actual, ('Actuals', 'B18'): 'Full-year actual'}, no_plan, 'measurement period')
    year = {**actual, ('Actuals', 'B18'): 'Full-year actual', ('Actuals', 'C18'): AS_OF - 365,
            (rd, 'G18'): AS_OF - 400, (rd, 'I18'): AS_OF - 366}
    case('Complete full-year actual qualifies', year, [(ck, 'L18', 1), (ck, 'AA18', 300)], 'measurement period')
    for cell, value in [('C18', AS_OF), ('D18', AS_OF + 1), ('K18', AS_OF - 2), ('K18', AS_OF + 1)]:
        case(f'Invalid measured chronology {cell} is excluded', {**actual, ('Actuals', cell): value}, no_plan, 'measurement period')
    case('Period ending before effective date cannot qualify', {**actual, ('Actuals', 'C18'): AS_OF - 50, ('Actuals', 'D18'): AS_OF - 40},
         no_plan + [(ck, 'I18', 0)], 'measurement period')
    case('Realized requires Implemented, not merely Approved', {**actual, (rd, 'E18'): 'Approved'},
         no_plan + [(ck, 'L18', 0)], 'measurement authorization')
    case('Confirmed zero actuals remain a measured result', {**actual, **{('Actuals', c + '18'): 0 for c in 'EFGH'}},
         [(ck, 'L18', 1), (ck, 'AA18', 0), (ck, 'AB18', 0), (ck, 'AC18', 0)], 'actual zero')
    case('Measured extra work and cash use count fully', {**actual, ('Actuals', 'E18'): -120, ('Actuals', 'F18'): -20,
         ('Actuals', 'G18'): -2000, (rd, 'J18'): 0, (rd, 'K18'): 0},
         [(ck, 'AA18', -120), (ck, 'AB18', -20), (ck, 'AC18', -2000)], 'signed actuals')
    case('Forecast rates do not rewrite stored actuals', {**actual, ('Settings', 'B9'): 120,
         ('Settings', 'B10'): 180, ('Settings', 'B11'): 150},
         [(ck, 'L18', 1), (ck, 'AA18', 300), (ck, 'AO18', 37000), ('Actuals', 'H18', 35000)], 'measurement independence')
    linked = demand_fixture(linked=True)
    case('Modeled case cannot remove linked demand', {**base, (rd, 'H18'): 'Modeled', **linked},
         [(ck, 'M18', 0), (re, 'L36', 'Count as extra demand'), (re, 'G11', 100), (re, 'E11', 0)], 'OC-05 overlap')
    case('Forecast and overlap use the same inclusion rule', {**base, **linked},
         [(ck, 'M18', 1), (re, 'L36', 'Included in eligible case'), (re, 'G11', 0), (re, 'E11', 400)], 'OC-05 overlap')
    case('Measured-only mode excludes forecast and restores demand', {**base, **linked, ('Settings', 'B8'): 'Measured only'},
         [(ck, 'K18', 1), (ck, 'M18', 0), (re, 'L36', 'Count as extra demand'), (re, 'G11', 100), (re, 'E11', 0)], 'OC-05 overlap')
    case('Measured-only mode includes actuals and excludes documented overlap', {**actual, **linked, ('Settings', 'B8'): 'Measured only'},
         [(ck, 'L18', 1), (ck, 'M18', 1), (re, 'L36', 'Included in eligible case'), (re, 'E11', 300)], 'OC-05 overlap')
    case('Blocked case restores independent demand', {**base, **linked, ('Challenge Review', 'J18'): 'Concern'},
         [(ck, 'M18', 0), (re, 'L36', 'Count as extra demand'), (re, 'G11', 100)], 'OC-05 overlap')
    case('Rejected case restores independent demand', {**base, **linked, (rd, 'E18'): 'Rejected'},
         [(ck, 'M18', 0), (re, 'L36', 'Count as extra demand')], 'OC-05 overlap')
    case('Example case cannot exclude linked Live demand', {**base, **linked, ('Cases', 'B18'): 'Example'},
         [(ck, 'M18', 0), (re, 'L36', 'Count as extra demand')], 'OC-05 overlap')
    case('Explicit extra demand is included beside eligible case', {**base, **demand_fixture()},
         [(re, 'L36', 'Count as extra demand'), (re, 'G11', 100), (re, 'I11', 300), (re, 'I14', 5000)], 'reinvestment demand')
    case('Unknown overlap key counts demand conservatively', {**base, **linked, (re, 'I36'): 'Unknown'},
         [(re, 'L36', 'Count as extra demand'), (re, 'G11', 100)], 'reinvestment demand')
    coverage_warning = 'Resolve incomplete demands or duplicate overlap claims before relying on coverage.'
    case('Duplicate overlap claims count both and request resolution', {**base, **linked, **demand_fixture(37, linked=True)},
         [(re, 'L36', 'Duplicate overlap — counted'), (re, 'L37', 'Duplicate overlap — counted'),
          (re, 'G11', 200), (re, 'A16', coverage_warning)], 'reinvestment demand')
    for col in 'CEFGHJK':
        case(f'Incomplete approved demand {col} stops coverage statement', {**base, **linked, (re, col + '36'): None},
             [(re, 'L36', 'Complete approved demand'), (re, 'A16', coverage_warning)], 'demand completeness')
    case('Negative demand cannot become a release', {**base, **linked, (re, 'E36'): -1},
         [(re, 'L36', 'Complete approved demand')], 'demand validation')
    case('Duplicate demand ID prevents clean coverage statement', {**base, **linked, (re, 'A36'): 'R001'},
         [(re, 'L36', 'Complete approved demand'), (re, 'A16', coverage_warning)], 'demand identity')
    warning = 'Integrity warning: missing readiness output or duplicate/missing key. Do not rely on planning totals.'
    case('Missing visible readiness output triggers integrity warning', {**base, (rd, 'M18'): None},
         [('Start', 'A12', warning)], 'OC-01 readiness integrity')
    case('Missing ledger readiness triggers integrity warning', {**base, (ck, 'J18'): None},
         [('Start', 'A12', warning)], 'OC-01 readiness integrity')
    case('Workload helper uses explicit per-person frequency', {('Settings', 'B17'): 100, ('Settings', 'B18'): 4,
         ('Settings', 'B19'): 15, ('Settings', 'B20'): 0, ('Settings', 'B21'): 0}, [('Settings', 'B22', 100)], 'workload helper')
    case('Workload helper rejects impossible follow-up share', {('Settings', 'B20'): 1.1}, [('Settings', 'B22', '')], 'workload helper')
    case('Workload helper intentionally does not change stored case hours', {('Settings', 'B17'): 4000},
         [('Settings', 'B22', 14666.666666666666), ('Current', 'C7', 7333.333333333333)], 'workload helper')
    all_scenarios = all(x['matches'] for c in scenarios for x in c['results'])
    success = not mismatches and not errors and all_scenarios and all(x['pass'] for x in structural + references)
    return {'review_date': '2026-10-23', 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'engine': 'Purpose-limited read-only source-formula interpreter; not native Microsoft Excel',
            'baseline': {'comparison': 'recalculated workbook' if recalculated else 'saved source caches',
                         'matched': matched, 'mismatches': mismatches, 'evaluation_errors': errors},
            'structural_checks': structural, 'independent_arithmetic': references, 'scenarios': scenarios,
            'summary': {'scenario_count': len(scenarios), 'scenario_assertions': sum(len(c['results']) for c in scenarios),
                        'arithmetic_assertions': len(references), 'structural_assertions': len(structural),
                        'all_observed_results_reproduced': all_scenarios, 'success': success}}


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('workbook', type=Path, nargs='?', default=Path(__file__).resolve().parents[1] / 'workbooks/Opportunity-Cost-Security-Model-v6.2.xlsx')
    p.add_argument('--recalculated', type=Path, help='Optional separately recalculated comparison file.')
    p.add_argument('--output', type=Path, required=True)
    args = p.parse_args()
    result = run(args.workbook, args.recalculated)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False, allow_nan=False) + '\n', encoding='utf-8')
    print(json.dumps({'baseline_matches': result['baseline']['matched'], 'baseline_mismatches': len(result['baseline']['mismatches']),
                      'evaluation_errors': len(result['baseline']['evaluation_errors']), **result['summary']}, indent=2))
    sys.exit(0 if result['summary']['success'] else 1)
