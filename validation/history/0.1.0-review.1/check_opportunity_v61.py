#!/usr/bin/env python3
"""Read-only source-formula scenarios; this is not an Excel engine."""
import argparse, json, math, sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_eval import Evaluator, XErr, num

class OpportunityEvaluator(Evaluator):
    def eval(self, node, sheet):
        if node[0] == 'fn' and node[1] == 'ROUND':
            x = num(self.eval(node[2][0], sheet)); n = int(num(self.eval(node[2][1], sheet)))
            return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))
        if node[0] == 'fn' and node[1] == 'HYPERLINK':
            return self.eval(node[2][-1], sheet)
        return super().eval(node, sheet)

def same(a,b):
    if a in (None,'') and b in (None,''): return True
    if isinstance(a,(int,float)) and isinstance(b,(int,float)):
        return math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-8)
    return a == b

def run(path, recalculated=None):
    ev=OpportunityEvaluator(path)
    wb=load_workbook(recalculated or path, read_only=True, data_only=True)
    cache={s.title:{c.coordinate:c.value for row in s for c in row if c.value is not None} for s in wb}; wb.close()
    matched=0; mismatches=[]; errors=[]
    for s,cells in ev.cells.items():
        for r,(value,_) in cells.items():
            if not (isinstance(value,str) and value.startswith('=')): continue
            try:
                got=ev.cell(s,r); expected=cache[s].get(r)
                if same(got,expected): matched+=1
                else: mismatches.append({'sheet':s,'cell':r,'expected':expected,'actual':got})
            except Exception as e: errors.append({'sheet':s,'cell':r,'error':str(e)})
    scenarios=[]
    def case(name, overrides, checks, finding=None):
        ev.reset(overrides); results=[]
        for s,r,expected in checks:
            try: actual=ev.cell(s,r)
            except Exception as e: actual=str(e)
            results.append({'sheet':s,'cell':r,'expected_observed_behavior':expected,'actual':actual,'matches':same(actual,expected)})
        scenarios.append({'name':name,'kind':'known-defect reproduction' if finding else 'behavior check','finding':finding,'results':results})
    rd='Review & Decision'; ca='Cases'; re='Results & Reinvestment'; cr='Challenge Review'
    case('Baseline examples are illustrative, no live cases',{},[('Start','C24',0),('Start','C37',1393050),('Start','C39',738250),(re,'C16',0)])
    case('Headcount driver recalculates workload',{('Parameters','C5'):4000},[(ca,'U6',14666.666666666666)])
    case('Annual hours mode ignores event headcount',{(ca,'N10'):9999},[(ca,'U10',100)])
    case('Unknown alternative licence blocks eligibility',{(ca,'AB6'):None},[(rd,'R6','unknown'),(rd,'AB6','no')])
    case('Confirmed zero alternative licence remains valid',{(ca,'AB6'):0},[(rd,'R6',1),(rd,'AB6','yes')])
    case('Explicit challenge blocker excludes the case',{(cr,'D24'):'blocked'},[(rd,'AA6','blocked'),(rd,'AB6','no')])
    case('Rejected case stays out of planning',{(rd,'Z6'):'rejected'},[(rd,'AB6','no')])
    case('Hypothesis cannot be relabelled validated',{(ca,'Z6'):'hypothesis',(ca,'AK6'):'validated'},[(rd,'AB6','no')])
    case('Missing team cost input still planning eligible',{(ca,'S6'):None,(ca,'AK6'):'validated'},[(ca,'W6',''),(rd,'R6',1),(rd,'AB6','yes'),(rd,'T6',-80)],'OC-03: Blank current team hours are costed as zero downstream.')
    case('Missing current cash input still planning eligible',{(ca,'T6'):None,(ca,'AK6'):'validated'},[(ca,'W6',''),(rd,'AB6','yes'),(rd,'U6',-60000)],'OC-03: Blank current recurring spend is costed as zero downstream.')
    case('Missing current person group still planning eligible',{(ca,'M6'):None,(ca,'AK6'):'validated'},[(rd,'J6',''),(rd,'AB6','yes')],'OC-03: Incomplete resource classification does not stop planning.')
    case('Realized without measurements: exclusion works, warning does not',{(ca,'AK6'):'realized'},[(rd,'AB6','no'),(rd,'AC6',None),('Start','C19','OK - realized entries are measurement-backed.')],'OC-01: Readiness AC6:AC20 is empty, so Start C19 is falsely reassuring.')
    case('Realized numbers without evidence or period accepted',{(ca,'AK6'):'realized',(ca,'AL6'):100,(ca,'AM6'):5,(ca,'AN6'):20},[(rd,'AB6','yes'),(re,'C16',100)],'OC-04: Numeric actuals pass without measurement period/evidence and Example data enters planning.')
    case('Example promoted to validated enters planning',{(ca,'AK6'):'validated'},[('Start','C24',0),(re,'C16',7033.333333333333)],'OC-04: Planning sums do not filter Case type Live.')
    case('Modeled case excludes reinvestment demand',{},[(re,'J5','no - priced in case 6'),(re,'C16',0)],'OC-05: Modeled/example replacement excludes demand before any validated/realized release is credited.')
    case('Zero scope remainder means full removal',{(ca,'AJ6'):'Reduce scope',(ca,'AI6'):0},[(rd,'R6',1)])
    case('Blank scope remainder means unknown',{(ca,'AJ6'):'Reduce scope',(ca,'AI6'):None},[(rd,'R6','unknown'),(rd,'AB6','no')])
    ev.reset()
    new={(ca,k+'18'):ev.cell(ca,k+'6') for k in ['C','D','E','F','G','H','I','J','L','M','N','O','P','Q','R','S','T','Y','Z','AA','AB','AC','AD','AE','AF','AH','AI','AJ','AK']}
    new.update({(ca,'B18'):13,(ca,'C18'):'Synthetic case 13',(ca,'D18'):'Live',(ca,'AK18'):'validated',(cr,'D36'):'blocked',('Business Model Matrix','D17'):4})
    # Active model is Automotive/TISAX; set its specific column as well.
    for col in ['D','E','F','G','H','I']: new[('Business Model Matrix',col+'17')]=4
    case('Spare row 18 blocker lost because review key is overwritten',new,[(rd,'AA18','unchecked'),(rd,'AB18','yes')],'OC-02: Challenge Review B36 is a status message, not the row-18 case ID.')
    return {'engine':'Purpose-limited read-only source-formula interpreter; not native Microsoft Excel',
            'baseline':{'matched':matched,'mismatches':mismatches,'evaluation_errors':errors},
            'scenarios':scenarios,'all_observed_results_reproduced':all(x['matches'] for c in scenarios for x in c['results'])}

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('workbook',type=Path); p.add_argument('--recalculated',type=Path);p.add_argument('--output',type=Path,required=True)
    args=p.parse_args();result=run(args.workbook,args.recalculated);args.output.write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps({'baseline':result['baseline'],'scenarios':len(result['scenarios']),'all_observed_results_reproduced':result['all_observed_results_reproduced']},indent=2))
    sys.exit(0 if not result['baseline']['mismatches'] and not result['baseline']['evaluation_errors'] and result['all_observed_results_reproduced'] else 1)
