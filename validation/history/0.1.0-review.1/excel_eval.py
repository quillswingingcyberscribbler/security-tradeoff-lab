"""Purpose-limited, read-only OOXML formula evaluator for audit scenarios.

This is an audit aid, not Excel. It evaluates the supported formulas directly
from a supplied workbook and keeps overrides in memory. No workbook is saved.
"""
import re,math,datetime
from functools import lru_cache
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries,get_column_letter
from openpyxl.utils.datetime import to_excel,from_excel

class XErr(Exception):pass
def flat(v):
 if isinstance(v,list):
  for x in v:yield from flat(x)
 else:yield v
def numeric(v):return isinstance(v,(float,int)) and not isinstance(v,bool)
def num(v):
 if v is None or v=='':return 0
 if numeric(v) or isinstance(v,bool):return float(v)
 try:return float(v)
 except:raise XErr('#VALUE!')
def text(v):
 if v is None:return ''
 if isinstance(v,bool):return 'TRUE' if v else 'FALSE'
 if numeric(v) and int(v)==v:return str(int(v))
 return str(v)
def truth(v):
 if isinstance(v,list):return any(truth(x) for x in flat(v))
 if v is None or v=='':return False
 if isinstance(v,str):
  if v.upper() in ['TRUE','FALSE']:return v.upper()=='TRUE'
  raise XErr('#VALUE!')
 return bool(v)
def map1(fn,v):return [[map1(fn,a) for a in row] for row in v] if isinstance(v,list) else fn(v)
def map2(fn,a,b):
 if not isinstance(a,list) and not isinstance(b,list):return fn(a,b)
 if isinstance(a,list) and isinstance(b,list):
  if len(a)!=len(b) or len(a[0])!=len(b[0]):raise XErr('#VALUE!')
  return [[map2(fn,x,y) for x,y in zip(ar,br)] for ar,br in zip(a,b)]
 if isinstance(a,list):return [[map2(fn,x,b) for x in ar] for ar in a]
 return [[map2(fn,a,y) for y in br] for br in b]
def cmp(a,b,op):
 if a is None:a='' if isinstance(b,str) else 0
 if b is None:b='' if isinstance(a,str) else 0
 if isinstance(a,str) and isinstance(b,str):a,b=a.lower(),b.lower()
 elif isinstance(a,str)!=isinstance(b,str):
  if op in ['=','<>']:return op=='<>'
  a,b=(1 if isinstance(a,str) else 0),(1 if isinstance(b,str) else 0)
 if op=='=':return a==b
 if op=='<>':return a!=b
 if op=='>':return a>b
 if op=='<':return a<b
 if op=='>=':return a>=b
 if op=='<=':return a<=b
def binary(op,a,b):
 def fn(x,y):
  if op in ['=','<>','>','<','>=','<=']:return cmp(x,y,op)
  if op=='&':return text(x)+text(y)
  x,y=num(x),num(y)
  if op=='+':return x+y
  if op=='-':return x-y
  if op=='*':return x*y
  if op=='/':
   if y==0:raise XErr('#DIV/0!')
   return x/y
  if op=='^':return x**y
  raise XErr('Unsupported operator '+op)
 return map2(fn,a,b)

class Parser:
 def __init__(self,s):self.ts=[t for t in Tokenizer(s).items if t.type!='WHITE-SPACE'];self.i=0
 def take(self):t=self.ts[self.i];self.i+=1;return t
 def peek(self):return self.ts[self.i] if self.i<len(self.ts) else None
 def expression(self,minp=0):
  t=self.take()
  if t.type=='OPERATOR-PREFIX':a=('unary',t.value,self.expression(6))
  elif t.type=='FUNC' and t.subtype=='OPEN':
   args=[]
   if self.peek().subtype!='CLOSE':
    while True:
     args.append(self.expression())
     if self.peek().type=='SEP':self.take();continue
     break
   self.take();a=('fn',t.value[:-1].replace('_xlfn.','').upper(),args)
  elif t.type=='PAREN' and t.subtype=='OPEN':a=self.expression();self.take()
  elif t.type=='OPERAND':
   if t.subtype=='NUMBER':a=('const',float(t.value))
   elif t.subtype=='TEXT':a=('const',t.value[1:-1].replace('""','"'))
   elif t.subtype=='LOGICAL':a=('const',t.value=='TRUE')
   elif t.subtype=='ERROR':a=('error',t.value)
   else:a=('ref',t.value)
  else:raise XErr('Unsupported token '+str(t))
  while self.peek() is not None:
   t=self.peek()
   if t.type=='OPERATOR-POSTFIX':self.take();a=('binary','/',a,('const',100));continue
   if t.type!='OPERATOR-INFIX':break
   p={'=':1,'<>':1,'<':1,'>':1,'<=':1,'>=':1,'&':2,'+':3,'-':3,'*':4,'/':4,'^':5}.get(t.value,-1)
   if p<minp:break
   self.take();a=('binary',t.value,a,self.expression(p+1))
  return a
@lru_cache(None)
def parse(s):
 p=Parser(s);result=p.expression()
 if p.i!=len(p.ts):raise XErr('Unparsed formula tokens: '+s)
 return result

class Evaluator:
 def __init__(self,path):
  wb=load_workbook(path,data_only=False,read_only=True)
  self.cells={s.title:{c.coordinate:(c.value,c.data_type) for row in s for c in row if c.value is not None} for s in wb}
  self.names={k:v.attr_text for k,v in wb.defined_names.items() if not k.startswith('_xlnm')}
  self.overrides={};self.memo={};self.active=set();wb.close()
 def reset(self,changes=None):self.overrides=dict(changes or {});self.memo={};self.active=set()
 def raw(self,s,a):return self.overrides.get((s,a),self.cells[s].get(a,(None,'n'))[0])
 def cell(self,s,a):
  a=a.replace('$','');k=(s,a)
  if k in self.memo:return self.memo[k]
  if k in self.active:raise XErr('Circular '+str(k))
  self.active.add(k)
  try:
   v=self.raw(s,a)
   if isinstance(v,str) and v.startswith('='):v=self.eval(parse(v),s)
   elif isinstance(v,datetime.datetime):v=to_excel(v)
   elif self.cells[s].get(a,(None,'n'))[1]=='e' and k not in self.overrides:raise XErr(v)
   self.memo[k]=v;return v
  finally:self.active.remove(k)
 def ref(self,r,s):
  if r in self.names:return self.ref(self.names[r],s)
  if '!' in r:s,r=r.rsplit('!',1);s=s.strip("'").replace("''", "'")
  r=r.replace('$','')
  if ':' not in r:return self.cell(s,r)
  c1,r1,c2,r2=range_boundaries(r)
  if r2 is None:r2=max(int(re.search(r'\d+',a).group()) for a in self.cells[s])
  if r1 is None:r1=1
  if c1 is None:c1=1
  if c2 is None:
   from openpyxl.utils.cell import column_index_from_string
   c2=max(column_index_from_string(re.match(r'[A-Z]+',a).group()) for a in self.cells[s])
  return [[self.cell(s,get_column_letter(c)+str(row)) for c in range(c1,c2+1)] for row in range(r1,r2+1)]
 def bounds(self,r,s):
  while r in self.names:r=self.names[r]
  if '!' in r:s,r=r.rsplit('!',1);s=s.strip("'").replace("''", "'")
  c1,r1,c2,r2=range_boundaries(r.replace('$',''))
  if r1 is None:r1=1
  if r2 is None:r2=max(int(re.search(r'\d+',a).group()) for a in self.cells[s])
  if c1 is None:c1=1
  if c2 is None:
   from openpyxl.utils.cell import column_index_from_string
   c2=max(column_index_from_string(re.match(r'[A-Z]+',a).group()) for a in self.cells[s])
  return s,c1,r1,c2,r2
 def eval(self,n,s):
  k=n[0]
  if k=='const':return n[1]
  if k=='error':raise XErr(n[1])
  if k=='ref':return self.ref(n[1],s)
  if k=='unary':return map1(lambda x:num(x)*(-1 if n[1]=='-' else 1),self.eval(n[2],s))
  if k=='binary':return binary(n[1],self.eval(n[2],s),self.eval(n[3],s))
  name,args=n[1:];ev=lambda i:self.eval(args[i],s)
  if name=='IF':return ev(1) if truth(ev(0)) else (ev(2) if len(args)>2 else False)
  if name in ['IFERROR','IFNA']:
   try:return ev(0)
   except XErr:return ev(1)
  if name in ['ISERROR','ISERR']:
   try:ev(0);return False
   except XErr:return True
  if name in ['ISNUMBER','ISTEXT','ISBLANK']:
   try:v=ev(0)
   except XErr:return False
   if name=='ISNUMBER':return map1(numeric,v)
   if name=='ISTEXT':return map1(lambda x:isinstance(x,str),v)
   return map1(lambda x:x is None,v)
  if name=='ISFORMULA':
   if args[0][0]!='ref':return False
   r=args[0][1];sn=s
   if '!' in r:sn,r=r.rsplit('!',1);sn=sn.strip("'")
   v=self.raw(sn,r.replace('$',''));return isinstance(v,str) and v.startswith('=')
  if name=='VLOOKUP' and args[1][0]=='ref':
   wanted=ev(0);col=int(num(ev(2)))
   sn,c1,r1,c2,r2=self.bounds(args[1][1],s)
   if len(args)>3 and truth(ev(3)):raise XErr('Unsupported approximate VLOOKUP')
   if col<1 or c1+col-1>c2:raise XErr('#REF!')
   for row in range(r1,r2+1):
    if cmp(wanted,self.cell(sn,get_column_letter(c1)+str(row)),'='):
     return self.cell(sn,get_column_letter(c1+col-1)+str(row))
   raise XErr('#N/A')
  if name=='INDEX' and args[0][0]=='ref':
   sn,c1,r1,c2,r2=self.bounds(args[0][1],s)
   row=int(num(ev(1)));col=int(num(ev(2))) if len(args)>2 else 1
   if r1==r2 and len(args)==2:col,row=row,1
   if row<1 or col<1 or r1+row-1>r2 or c1+col-1>c2:raise XErr('#REF!')
   return self.cell(sn,get_column_letter(c1+col-1)+str(r1+row-1))
  values=[ev(i) for i in range(len(args))]
  vals=list(flat(values))
  if name=='TRUE':return True
  if name=='FALSE':return False
  if name=='NOT':return not truth(values[0])
  if name=='AND':return all(truth(x) for x in vals)
  if name=='OR':return any(truth(x) for x in vals)
  if name=='ISNUMBER':return map1(numeric,values[0])
  if name=='ISTEXT':return map1(lambda x:isinstance(x,str),values[0])
  if name=='ISBLANK':return map1(lambda x:x is None,values[0])
  if name=='INT':return map1(lambda x:math.floor(num(x)),values[0])
  if name=='ABS':return map1(lambda x:abs(num(x)),values[0])
  if name in ['SUM','MIN','MAX','COUNT','AVERAGE']:
   nums=[x for x in vals if numeric(x)]
   if name=='COUNT':return len(nums)
   if name=='SUM':return sum(nums)
   if name=='MIN':return min(nums,default=0)
   if name=='MAX':return max(nums,default=0)
   return sum(nums)/len(nums) if nums else 0
  if name=='COUNTA':return sum(x is not None for x in vals)
  if name=='COUNTBLANK':return sum(x is None or x=='' for x in vals)
  if name=='SUMPRODUCT':
   a=values[0]
   for b in values[1:]:a=map2(lambda x,y:num(x)*num(y),a,b)
   return sum(num(x) for x in flat(a))
  if name in ['COUNTIF','COUNTIFS','SUMIF','SUMIFS']:
   if name=='SUMIFS':sumrange=list(flat(values[0]));pairs=values[1:]
   elif name=='SUMIF':sumrange=list(flat(values[2] if len(values)>2 else values[0]));pairs=values[:2]
   else:sumrange=None;pairs=values
   ranges=[list(flat(pairs[i])) for i in range(0,len(pairs),2)];crit=pairs[1::2]
   good=[all(self.criterion(r[j],c) for r,c in zip(ranges,crit)) for j in range(len(ranges[0]))]
   return sum(num(sumrange[j]) for j,g in enumerate(good) if g) if sumrange is not None else sum(good)
  if name in ['LEFT','RIGHT']:
   st=text(values[0]);ct=int(num(values[1])) if len(values)>1 else 1
   return st[:ct] if name=='LEFT' else st[len(st)-ct:]
  if name=='MID':return text(values[0])[int(num(values[1]))-1:int(num(values[1]))-1+int(num(values[2]))]
  if name=='LEN':return len(text(values[0]))
  if name=='TEXTJOIN':return text(values[0]).join(text(x) for x in flat(values[2:]) if not truth(values[1]) or text(x)!='')
  if name=='TEXT':
   v,f=values
   if f=='00':return f'{int(num(v)):02}'
   if f=='0%':return str(round(num(v)*100))+'%'
   return text(v)
  if name=='DATE':
   y,m,d=[int(num(x)) for x in values];y=y+1900 if 0<=y<1900 else y
   y+=(m-1)//12;m=(m-1)%12+1
   try:return to_excel(datetime.datetime(y,m,1)+datetime.timedelta(days=d-1))
   except:raise XErr('#NUM!')
  if name in ['YEAR','MONTH','DAY']:
   dt=from_excel(num(values[0]));return getattr(dt,name.lower())
  if name=='MATCH':
   wanted,arr=values[:2]
   for i,v in enumerate(flat(arr)):
    if cmp(wanted,v,'='):return i+1
   raise XErr('#N/A')
  if name=='INDEX':
   arr=values[0];row=int(num(values[1]));col=int(num(values[2])) if len(values)>2 else 1
   if not isinstance(arr,list):return arr
   if len(arr)==1 and len(values)==2:return arr[0][row-1]
   return arr[row-1][col-1]
  if name=='VLOOKUP':
   wanted,arr,col=values[:3]
   for row in arr:
    if cmp(wanted,row[0],'='):return row[int(num(col))-1]
   raise XErr('#N/A')
  if name=='ROW':
   return int(re.search(r'\d+',args[0][1]).group())
  raise XErr('Unsupported function '+name)
 def criterion(self,v,c):
  if not isinstance(c,str):return cmp(v,c,'=')
  m=re.match(r'^(>=|<=|<>|>|<|=)?(.*)$',c);op,rhs=m.groups();op=op or '='
  if '*' in rhs or '?' in rhs:
   pattern='^'+re.escape(rhs).replace(r'\*','.*').replace(r'\?','.')+'$'
   ok=bool(re.match(pattern,text(v),re.I));return not ok if op=='<>' else ok
  try:rhs=float(rhs)
  except ValueError:pass
  return cmp(v,rhs,op)
